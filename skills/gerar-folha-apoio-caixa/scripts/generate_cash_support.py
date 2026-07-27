#!/usr/bin/env python3
"""Generate Hotel Oslo daily cash-support sheets from normalized JSON.

Input is intentionally independent of the InvoiceXpress MCP tool names. The
agent queries the MCP, applies the business rules, and supplies only the
normalized movements described in the skill.
"""

from __future__ import annotations

import argparse
import json
import math
import re
import sys
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Any

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.pagebreak import Break
except ModuleNotFoundError as exc:  # pragma: no cover - environment dependent
    raise SystemExit(
        "Dependência em falta: openpyxl. Instale com `python -m pip install openpyxl`."
    ) from exc

try:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Table, TableStyle
except ModuleNotFoundError as exc:  # pragma: no cover - environment dependent
    raise SystemExit(
        "Dependência em falta: reportlab. Instale com `python -m pip install reportlab`."
    ) from exc


NAVY = "173F59"
SLATE = "5D7687"
PALE_YELLOW = "FFF2CC"
PALE_BLUE = "DDEBF7"
PALE_RED = "FCE4E4"
BORDER_COLOR = "B8C9D6"
DARK_TEXT = "1F2937"
RETURN_RED = "D7191C"
WHITE = "FFFFFF"

MIN_MOVEMENT_SLOTS = 17
CURRENCY_QUANTUM = Decimal("0.01")
VALID_MOVEMENT_TYPES = {"InvoicePayment", "InvoiceReturn"}
VALID_SOURCE_TYPES = {"Invoice", "CreditNote"}
SAFE_BASENAME = re.compile(r"^[A-Za-z0-9À-ÿ_.-]+$")


class InputError(ValueError):
    """Raised when normalized input violates the skill contract."""


@dataclass(frozen=True)
class Movement:
    document: str
    amount: Decimal
    movement_type: str
    source_type: str
    source_id: str | None


@dataclass(frozen=True)
class DayData:
    day: date
    provisional: bool
    movements: tuple[Movement, ...]

    @property
    def payment_total(self) -> Decimal:
        return money(
            sum(
                (movement.amount for movement in self.movements
                 if movement.movement_type == "InvoicePayment"),
                Decimal("0"),
            )
        )

    @property
    def return_total(self) -> Decimal:
        return money(
            sum(
                (movement.amount for movement in self.movements
                 if movement.movement_type == "InvoiceReturn"),
                Decimal("0"),
            )
        )

    @property
    def net_total(self) -> Decimal:
        return money(self.payment_total + self.return_total)


def money(value: Decimal) -> Decimal:
    return value.quantize(CURRENCY_QUANTUM, rounding=ROUND_HALF_UP)


def parse_decimal(value: Any, label: str) -> Decimal:
    if isinstance(value, bool) or value is None:
        raise InputError(f"{label} tem de ser um número.")
    try:
        parsed = Decimal(str(value))
    except (InvalidOperation, ValueError) as exc:
        raise InputError(f"{label} tem um valor inválido: {value!r}.") from exc
    if not parsed.is_finite():
        raise InputError(f"{label} tem de ser finito.")
    return money(parsed)


def parse_iso_date(value: Any, label: str) -> date:
    if not isinstance(value, str):
        raise InputError(f"{label} tem de usar o formato YYYY-MM-DD.")
    try:
        return date.fromisoformat(value)
    except ValueError as exc:
        raise InputError(f"{label} inválida: {value!r}.") from exc


def require_text(value: Any, label: str) -> str:
    if not isinstance(value, str) or not value.strip():
        raise InputError(f"{label} tem de ser texto não vazio.")
    return value.strip()


def load_days(input_path: Path) -> tuple[DayData, ...]:
    try:
        payload = json.loads(input_path.read_text(encoding="utf-8"))
    except FileNotFoundError as exc:
        raise InputError(f"Ficheiro não encontrado: {input_path}") from exc
    except json.JSONDecodeError as exc:
        raise InputError(f"JSON inválido em {input_path}: {exc}") from exc

    if not isinstance(payload, dict):
        raise InputError("A raiz do JSON tem de ser um objeto.")
    if payload.get("timezone", "Europe/Lisbon") != "Europe/Lisbon":
        raise InputError("A timezone tem de ser Europe/Lisbon.")

    raw_days = payload.get("days")
    if not isinstance(raw_days, list) or not raw_days:
        raise InputError("O JSON tem de conter uma lista `days` não vazia.")

    seen_dates: set[date] = set()
    parsed_days: list[DayData] = []

    for day_index, raw_day in enumerate(raw_days):
        if not isinstance(raw_day, dict):
            raise InputError(f"days[{day_index}] tem de ser um objeto.")
        parsed_date = parse_iso_date(raw_day.get("date"), f"days[{day_index}].date")
        if parsed_date in seen_dates:
            raise InputError(f"Data duplicada no JSON: {parsed_date.isoformat()}.")
        seen_dates.add(parsed_date)

        provisional = raw_day.get("provisional", False)
        if not isinstance(provisional, bool):
            raise InputError(f"days[{day_index}].provisional tem de ser booleano.")

        raw_movements = raw_day.get("movements")
        if not isinstance(raw_movements, list):
            raise InputError(f"days[{day_index}].movements tem de ser uma lista.")

        seen_documents: set[str] = set()
        seen_source_ids: set[str] = set()
        movements: list[Movement] = []

        for movement_index, raw_movement in enumerate(raw_movements):
            prefix = f"days[{day_index}].movements[{movement_index}]"
            if not isinstance(raw_movement, dict):
                raise InputError(f"{prefix} tem de ser um objeto.")

            document = require_text(raw_movement.get("document"), f"{prefix}.document")
            if document in seen_documents:
                raise InputError(
                    f"Documento duplicado em {parsed_date.isoformat()}: {document}."
                )
            seen_documents.add(document)

            movement_type = require_text(
                raw_movement.get("movement_type"), f"{prefix}.movement_type"
            )
            if movement_type not in VALID_MOVEMENT_TYPES:
                raise InputError(
                    f"{prefix}.movement_type só pode ser InvoicePayment ou InvoiceReturn."
                )

            source_type = require_text(
                raw_movement.get("source_type"), f"{prefix}.source_type"
            )
            if source_type not in VALID_SOURCE_TYPES:
                raise InputError(
                    f"{prefix}.source_type só pode ser Invoice ou CreditNote."
                )
            if (
                movement_type == "InvoicePayment" and source_type != "Invoice"
            ) or (
                movement_type == "InvoiceReturn" and source_type != "CreditNote"
            ):
                raise InputError(
                    f"{prefix}: movimento {movement_type} incompatível com {source_type}."
                )

            amount = parse_decimal(raw_movement.get("amount"), f"{prefix}.amount")
            if movement_type == "InvoicePayment" and amount <= 0:
                raise InputError(f"{prefix}: InvoicePayment tem de ser positivo.")
            if movement_type == "InvoiceReturn" and amount >= 0:
                raise InputError(f"{prefix}: InvoiceReturn tem de ser negativo.")

            source_id_value = raw_movement.get("source_id")
            source_id = None
            if source_id_value is not None:
                source_id = require_text(source_id_value, f"{prefix}.source_id")
                if source_id in seen_source_ids:
                    raise InputError(
                        f"source_id duplicado em {parsed_date.isoformat()}: {source_id}."
                    )
                seen_source_ids.add(source_id)

            movements.append(
                Movement(
                    document=document,
                    amount=amount,
                    movement_type=movement_type,
                    source_type=source_type,
                    source_id=source_id,
                )
            )

        parsed_days.append(
            DayData(
                day=parsed_date,
                provisional=provisional,
                movements=tuple(movements),
            )
        )

    return tuple(sorted(parsed_days, key=lambda item: item.day))


def default_basename(days: tuple[DayData, ...]) -> str:
    if len(days) == 1:
        return f"Folha_apoio_caixa_{days[0].day.isoformat()}"
    return (
        f"Folhas_apoio_caixa_{days[0].day.isoformat()}"
        f"_a_{days[-1].day.isoformat()}"
    )


def validate_basename(value: str) -> str:
    candidate = Path(value).name
    if candidate != value or not SAFE_BASENAME.fullmatch(candidate):
        raise InputError(
            "O basename só pode conter letras, números, ponto, hífen e underscore."
        )
    for suffix in (".xlsx", ".pdf"):
        if candidate.lower().endswith(suffix):
            candidate = candidate[: -len(suffix)]
    if not candidate:
        raise InputError("O basename não pode ficar vazio.")
    return candidate


def excel_title(day_data: DayData) -> str:
    text = day_data.day.strftime("%d/%m/%Y")
    if day_data.provisional:
        text += " · PROVISÓRIA - DIA EM CURSO"
    return text


def movement_capacity(count: int) -> int:
    return max(MIN_MOVEMENT_SLOTS, math.ceil(max(count, 1) / 17) * 17)


def create_excel(days: tuple[DayData, ...], output_path: Path) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)

    thin_side = Side(style="thin", color=BORDER_COLOR)
    medium_navy = Side(style="medium", color=NAVY)
    double_navy = Side(style="double", color=NAVY)
    all_border = Border(
        left=thin_side,
        right=thin_side,
        top=thin_side,
        bottom=thin_side,
    )

    for day_data in days:
        sheet_name = day_data.day.strftime("%d-%m-%Y")
        sheet = workbook.create_sheet(sheet_name)
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = "A6"

        capacity = movement_capacity(len(day_data.movements))
        movement_start = 6
        movement_end = movement_start + capacity - 1
        total_start = movement_end + 1
        total_end = total_start + 2
        summary_title_row = total_end + 2
        summary_header_row = summary_title_row + 1
        summary_first_row = summary_header_row + 1
        summary_last_row = summary_first_row + 3

        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=7)
        sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=7)
        sheet.merge_cells(start_row=3, start_column=1, end_row=3, end_column=7)
        for row in range(total_start, total_end + 1):
            sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        sheet.merge_cells(
            start_row=summary_title_row,
            start_column=1,
            end_row=summary_title_row,
            end_column=7,
        )
        sheet.merge_cells(
            start_row=summary_header_row,
            start_column=1,
            end_row=summary_header_row,
            end_column=6,
        )
        for row in range(summary_first_row, summary_last_row + 1):
            sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)

        sheet["A1"] = "FOLHA DE APOIO À CAIXA"
        sheet["A2"] = excel_title(day_data)
        sheet["A3"] = "Preencher manualmente os campos assinalados."
        headers = [
            "DOCUMENTO",
            "FOLIO",
            "VALOR",
            "TIPO",
            "HORA",
            "DIA",
            "MEIO DE PAGAMENTO",
        ]
        for column, value in enumerate(headers, start=1):
            sheet.cell(row=5, column=column, value=value)

        title_cell = sheet["A1"]
        title_cell.fill = PatternFill("solid", fgColor=NAVY)
        title_cell.font = Font(name="Arial", size=22, bold=True, color=WHITE)
        title_cell.alignment = Alignment(horizontal="center", vertical="center")

        date_cell = sheet["A2"]
        date_cell.fill = PatternFill("solid", fgColor=PALE_YELLOW)
        date_cell.font = Font(name="Arial", size=14, bold=True, color="6B4E00")
        date_cell.alignment = Alignment(horizontal="center", vertical="center")

        note_cell = sheet["A3"]
        note_cell.font = Font(
            name="Arial", size=10, italic=True, color="64748B"
        )
        note_cell.alignment = Alignment(horizontal="left", vertical="center")

        for cell in sheet[5]:
            cell.fill = PatternFill("solid", fgColor=SLATE)
            cell.font = Font(name="Arial", size=10, bold=True, color=WHITE)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = all_border

        for offset in range(capacity):
            row = movement_start + offset
            movement = (
                day_data.movements[offset]
                if offset < len(day_data.movements)
                else None
            )

            if movement:
                values: list[Any] = [
                    movement.document,
                    "____________",
                    float(movement.amount),
                    movement.movement_type,
                    "____:____",
                    datetime.combine(day_data.day, datetime.min.time()),
                    "________________",
                ]
            else:
                values = [
                    "____________",
                    "____________",
                    "________ €",
                    "____________",
                    "____:____",
                    datetime.combine(day_data.day, datetime.min.time()),
                    "________________",
                ]

            for column, value in enumerate(values, start=1):
                cell = sheet.cell(row=row, column=column, value=value)
                cell.font = Font(name="Arial", size=9, color=DARK_TEXT)
                cell.border = all_border
                cell.alignment = Alignment(vertical="center")

            sheet.cell(row=row, column=1).alignment = Alignment(
                horizontal="left", vertical="center"
            )
            sheet.cell(row=row, column=2).alignment = Alignment(
                horizontal="left", vertical="center"
            )
            sheet.cell(row=row, column=3).alignment = Alignment(
                horizontal="right", vertical="center"
            )
            for column in (4, 5, 6):
                sheet.cell(row=row, column=column).alignment = Alignment(
                    horizontal="center", vertical="center"
                )
            sheet.cell(row=row, column=7).alignment = Alignment(
                horizontal="left", vertical="center"
            )
            sheet.cell(row=row, column=6).number_format = "dd/mm/yyyy"

            if movement:
                sheet.cell(row=row, column=3).number_format = '#,##0.00 "€"'
                for column in (2, 5, 7):
                    sheet.cell(row=row, column=column).fill = PatternFill(
                        "solid", fgColor=PALE_YELLOW
                    )
                if movement.movement_type == "InvoiceReturn":
                    for column in (1, 3, 4, 6):
                        sheet.cell(row=row, column=column).fill = PatternFill(
                            "solid", fgColor=PALE_RED
                        )
                    for column in (3, 4):
                        sheet.cell(row=row, column=column).font = Font(
                            name="Arial",
                            size=9,
                            bold=True,
                            color=RETURN_RED,
                        )
            else:
                for column in (1, 2, 3, 4, 5, 7):
                    sheet.cell(row=row, column=column).fill = PatternFill(
                        "solid", fgColor=PALE_YELLOW
                    )
                sheet.cell(row=row, column=6).fill = PatternFill(
                    "solid", fgColor=PALE_BLUE
                )

        total_labels = ["TOTAL PAGAMENTOS", "TOTAL DEVOLUÇÕES", "TOTAL LÍQUIDO"]
        for offset, label in enumerate(total_labels):
            row = total_start + offset
            sheet.cell(row=row, column=1, value=label)
            for column in range(1, 4):
                cell = sheet.cell(row=row, column=column)
                cell.fill = PatternFill("solid", fgColor=PALE_BLUE)
                cell.font = Font(name="Arial", size=10, bold=True, color=NAVY)
                cell.border = all_border
                cell.alignment = Alignment(horizontal="right", vertical="center")
            sheet.cell(row=row, column=3).number_format = '#,##0.00 "€"'

        sheet.cell(
            row=total_start,
            column=3,
            value=(
                f'=SUMIF(D{movement_start}:D{movement_end},'
                f'"InvoicePayment",C{movement_start}:C{movement_end})'
            ),
        )
        sheet.cell(
            row=total_start + 1,
            column=3,
            value=(
                f'=SUMIF(D{movement_start}:D{movement_end},'
                f'"InvoiceReturn",C{movement_start}:C{movement_end})'
            ),
        )
        sheet.cell(
            row=total_end,
            column=3,
            value=f"=C{total_start}+C{total_start + 1}",
        )
        sheet.cell(row=total_start + 1, column=3).font = Font(
            name="Arial", size=10, bold=True, color=RETURN_RED
        )
        for column in range(1, 4):
            sheet.cell(row=total_end, column=column).border = Border(
                left=thin_side,
                right=thin_side,
                top=double_navy,
                bottom=medium_navy,
            )

        sheet.cell(
            row=summary_title_row,
            column=1,
            value="RESUMO POR MEIO DE PAGAMENTO",
        )
        sheet.cell(row=summary_header_row, column=1, value="MEIO DE PAGAMENTO")
        sheet.cell(row=summary_header_row, column=7, value="TOTAL")
        payment_methods = [
            ("Cartão de Crédito", "________________ €"),
            ("Dinheiro", "________________ €"),
            ("Transferência", "________________ €"),
            ("Current Account", "Anexar Folha do Programa"),
        ]
        for offset, (method, value) in enumerate(payment_methods):
            row = summary_first_row + offset
            sheet.cell(row=row, column=1, value=method)
            sheet.cell(row=row, column=7, value=value)

        for column in range(1, 8):
            cell = sheet.cell(row=summary_title_row, column=column)
            cell.fill = PatternFill("solid", fgColor=NAVY)
            cell.font = Font(name="Arial", size=12, bold=True, color=WHITE)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = all_border

            cell = sheet.cell(row=summary_header_row, column=column)
            cell.fill = PatternFill("solid", fgColor=SLATE)
            cell.font = Font(name="Arial", size=10, bold=True, color=WHITE)
            cell.alignment = Alignment(
                horizontal="right" if column == 7 else "left",
                vertical="center",
            )
            cell.border = all_border

        for row in range(summary_first_row, summary_last_row + 1):
            for column in range(1, 8):
                cell = sheet.cell(row=row, column=column)
                cell.font = Font(name="Arial", size=10, color=DARK_TEXT)
                cell.alignment = Alignment(
                    horizontal="right" if column == 7 else "left",
                    vertical="center",
                )
                cell.border = all_border
            sheet.cell(row=row, column=7).fill = PatternFill(
                "solid", fgColor=PALE_YELLOW
            )
        sheet.cell(row=summary_last_row, column=7).font = Font(
            name="Arial", size=8, bold=True, color=DARK_TEXT
        )
        sheet.cell(row=summary_last_row, column=7).alignment = Alignment(
            horizontal="center", vertical="center"
        )

        widths = {
            "A": 18,
            "B": 15,
            "C": 13,
            "D": 18,
            "E": 11,
            "F": 14,
            "G": 23,
        }
        for column, width in widths.items():
            sheet.column_dimensions[column].width = width

        sheet.row_dimensions[1].height = 42
        sheet.row_dimensions[2].height = 30
        sheet.row_dimensions[3].height = 24
        sheet.row_dimensions[4].height = 8
        sheet.row_dimensions[5].height = 32
        for row in range(movement_start, movement_end + 1):
            sheet.row_dimensions[row].height = 24
        for row in range(total_start, total_end + 1):
            sheet.row_dimensions[row].height = 22
        sheet.row_dimensions[summary_title_row].height = 27
        sheet.row_dimensions[summary_header_row].height = 24
        for row in range(summary_first_row, summary_last_row + 1):
            sheet.row_dimensions[row].height = 23

        sheet.page_setup.orientation = "portrait"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.page_margins.left = 0.25
        sheet.page_margins.right = 0.25
        sheet.page_margins.top = 0.35
        sheet.page_margins.bottom = 0.35
        sheet.print_options.horizontalCentered = True
        sheet.print_title_rows = "1:5"
        sheet.print_area = f"A1:G{summary_last_row}"

        if capacity > MIN_MOVEMENT_SLOTS:
            for break_row in range(
                movement_start + MIN_MOVEMENT_SLOTS - 1,
                movement_end,
                MIN_MOVEMENT_SLOTS,
            ):
                sheet.row_breaks.append(Break(id=break_row))

    if hasattr(workbook, "calculation"):
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
        workbook.calculation.calcMode = "auto"

    workbook.save(output_path)
    verify_excel_structure(output_path, days)


def verify_excel_structure(output_path: Path, days: tuple[DayData, ...]) -> None:
    workbook = load_workbook(output_path, data_only=False, read_only=False)
    expected_sheets = [day_data.day.strftime("%d-%m-%Y") for day_data in days]
    if workbook.sheetnames != expected_sheets:
        raise RuntimeError(
            f"Worksheets inesperadas: {workbook.sheetnames!r}; "
            f"esperado {expected_sheets!r}."
        )

    for day_data in days:
        sheet = workbook[day_data.day.strftime("%d-%m-%Y")]
        if sheet["B5"].value != "FOLIO":
            raise RuntimeError("O cabeçalho FOLIO não foi gravado corretamente.")
        documents = [
            sheet.cell(row=6 + index, column=1).value
            for index in range(len(day_data.movements))
        ]
        expected_documents = [
            movement.document for movement in day_data.movements
        ]
        if documents != expected_documents:
            raise RuntimeError(
                f"Documentos incorretos na folha {sheet.title}: {documents!r}."
            )

        capacity = movement_capacity(len(day_data.movements))
        total_start = 6 + capacity
        summary_last_row = total_start + 9
        if sheet.cell(row=summary_last_row, column=7).value != (
            "Anexar Folha do Programa"
        ):
            raise RuntimeError(
                f"Current Account incorreto na folha {sheet.title}."
            )
        if sheet.page_setup.orientation != "portrait":
            raise RuntimeError(f"A folha {sheet.title} não está em orientação vertical.")


def register_pdf_fonts() -> tuple[str, str]:
    candidates = [
        (
            Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
            Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
        ),
        (
            Path("/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf"),
            Path("/usr/share/fonts/truetype/liberation2/LiberationSans-Bold.ttf"),
        ),
    ]
    for regular, bold in candidates:
        if regular.exists() and bold.exists():
            pdfmetrics.registerFont(TTFont("CashSans", str(regular)))
            pdfmetrics.registerFont(TTFont("CashSansBold", str(bold)))
            return "CashSans", "CashSansBold"
    return "Helvetica", "Helvetica-Bold"


def pdf_amount(value: Decimal) -> str:
    return f"{value:,.2f} €"


def p(text: Any, style: ParagraphStyle) -> Paragraph:
    return Paragraph(str(text), style)


def create_pdf(days: tuple[DayData, ...], output_path: Path) -> None:
    regular_font, bold_font = register_pdf_fonts()
    navy = colors.HexColor(f"#{NAVY}")
    slate = colors.HexColor(f"#{SLATE}")
    pale_yellow = colors.HexColor(f"#{PALE_YELLOW}")
    pale_blue = colors.HexColor(f"#{PALE_BLUE}")
    pale_red = colors.HexColor(f"#{PALE_RED}")
    border = colors.HexColor(f"#{BORDER_COLOR}")
    dark_text = colors.HexColor(f"#{DARK_TEXT}")
    return_red = colors.HexColor(f"#{RETURN_RED}")
    white = colors.white

    styles = {
        "title": ParagraphStyle(
            "title",
            fontName=bold_font,
            fontSize=20,
            leading=23,
            textColor=white,
            alignment=TA_CENTER,
        ),
        "date": ParagraphStyle(
            "date",
            fontName=bold_font,
            fontSize=12,
            leading=15,
            textColor=colors.HexColor("#6B4E00"),
            alignment=TA_CENTER,
        ),
        "note": ParagraphStyle(
            "note",
            fontName=regular_font,
            fontSize=8.5,
            leading=10,
            textColor=colors.HexColor("#64748B"),
            alignment=TA_LEFT,
        ),
        "header": ParagraphStyle(
            "header",
            fontName=bold_font,
            fontSize=7.6,
            leading=9,
            textColor=white,
            alignment=TA_CENTER,
        ),
        "header_left": ParagraphStyle(
            "header_left",
            fontName=bold_font,
            fontSize=7.6,
            leading=9,
            textColor=white,
            alignment=TA_LEFT,
        ),
        "header_right": ParagraphStyle(
            "header_right",
            fontName=bold_font,
            fontSize=7.6,
            leading=9,
            textColor=white,
            alignment=TA_RIGHT,
        ),
        "body_left": ParagraphStyle(
            "body_left",
            fontName=regular_font,
            fontSize=7.3,
            leading=8.5,
            textColor=dark_text,
            alignment=TA_LEFT,
        ),
        "body_center": ParagraphStyle(
            "body_center",
            fontName=regular_font,
            fontSize=7.3,
            leading=8.5,
            textColor=dark_text,
            alignment=TA_CENTER,
        ),
        "body_right": ParagraphStyle(
            "body_right",
            fontName=regular_font,
            fontSize=7.3,
            leading=8.5,
            textColor=dark_text,
            alignment=TA_RIGHT,
        ),
        "return_center": ParagraphStyle(
            "return_center",
            fontName=bold_font,
            fontSize=7.3,
            leading=8.5,
            textColor=return_red,
            alignment=TA_CENTER,
        ),
        "return_right": ParagraphStyle(
            "return_right",
            fontName=bold_font,
            fontSize=7.3,
            leading=8.5,
            textColor=return_red,
            alignment=TA_RIGHT,
        ),
        "total_label": ParagraphStyle(
            "total_label",
            fontName=bold_font,
            fontSize=8.2,
            leading=9.5,
            textColor=navy,
            alignment=TA_RIGHT,
        ),
        "total_value": ParagraphStyle(
            "total_value",
            fontName=bold_font,
            fontSize=8.2,
            leading=9.5,
            textColor=navy,
            alignment=TA_RIGHT,
        ),
        "total_return": ParagraphStyle(
            "total_return",
            fontName=bold_font,
            fontSize=8.2,
            leading=9.5,
            textColor=return_red,
            alignment=TA_RIGHT,
        ),
        "summary_title": ParagraphStyle(
            "summary_title",
            fontName=bold_font,
            fontSize=10,
            leading=12,
            textColor=white,
            alignment=TA_CENTER,
        ),
        "current_account": ParagraphStyle(
            "current_account",
            fontName=bold_font,
            fontSize=5.8,
            leading=6.8,
            textColor=dark_text,
            alignment=TA_CENTER,
        ),
    }

    page_width, _ = A4
    usable_width = page_width - 20 * mm
    column_widths = [
        87 * usable_width / 539,
        74 * usable_width / 539,
        61 * usable_width / 539,
        88 * usable_width / 539,
        54 * usable_width / 539,
        70 * usable_width / 539,
        105 * usable_width / 539,
    ]

    document = SimpleDocTemplate(
        str(output_path),
        pagesize=A4,
        leftMargin=10 * mm,
        rightMargin=10 * mm,
        topMargin=8 * mm,
        bottomMargin=8 * mm,
        title="Folha de Apoio à Caixa",
        author="Hotel Oslo Coimbra",
    )
    story: list[Any] = []

    pages: list[tuple[DayData, list[Movement], int, int]] = []
    for day_data in days:
        chunks = [
            list(day_data.movements[index:index + MIN_MOVEMENT_SLOTS])
            for index in range(0, len(day_data.movements), MIN_MOVEMENT_SLOTS)
        ] or [[]]
        total_pages = len(chunks)
        for page_index, chunk in enumerate(chunks, start=1):
            pages.append((day_data, chunk, page_index, total_pages))

    for global_index, (day_data, chunk, page_index, total_pages) in enumerate(pages):
        is_final_day_page = page_index == total_pages
        date_text = excel_title(day_data)
        if total_pages > 1:
            date_text += f" · PÁGINA {page_index}/{total_pages}"

        rows: list[list[Any]] = [
            [p("FOLHA DE APOIO À CAIXA", styles["title"]), "", "", "", "", "", ""],
            [p(date_text, styles["date"]), "", "", "", "", "", ""],
            [
                p(
                    "Preencher manualmente os campos assinalados.",
                    styles["note"],
                ),
                "",
                "",
                "",
                "",
                "",
                "",
            ],
            ["", "", "", "", "", "", ""],
            [
                p("DOCUMENTO", styles["header"]),
                p("FOLIO", styles["header"]),
                p("VALOR", styles["header"]),
                p("TIPO", styles["header"]),
                p("HORA", styles["header"]),
                p("DIA", styles["header"]),
                p("MEIO DE PAGAMENTO", styles["header"]),
            ],
        ]

        for slot in range(MIN_MOVEMENT_SLOTS):
            movement = chunk[slot] if slot < len(chunk) else None
            if movement:
                amount_style = (
                    styles["return_right"]
                    if movement.movement_type == "InvoiceReturn"
                    else styles["body_right"]
                )
                type_style = (
                    styles["return_center"]
                    if movement.movement_type == "InvoiceReturn"
                    else styles["body_center"]
                )
                rows.append(
                    [
                        p(movement.document, styles["body_left"]),
                        p("____________", styles["body_left"]),
                        p(pdf_amount(movement.amount), amount_style),
                        p(movement.movement_type, type_style),
                        p("____:____", styles["body_center"]),
                        p(day_data.day.strftime("%d/%m/%Y"), styles["body_center"]),
                        p("________________", styles["body_left"]),
                    ]
                )
            else:
                rows.append(
                    [
                        p("____________", styles["body_left"]),
                        p("____________", styles["body_left"]),
                        p("________ €", styles["body_right"]),
                        p("____________", styles["body_center"]),
                        p("____:____", styles["body_center"]),
                        p(day_data.day.strftime("%d/%m/%Y"), styles["body_center"]),
                        p("________________", styles["body_left"]),
                    ]
                )

        total_row_start: int | None = None
        summary_title_row: int | None = None
        summary_header_row: int | None = None
        summary_first_row: int | None = None

        if is_final_day_page:
            total_row_start = len(rows)
            rows.extend(
                [
                    [
                        p("TOTAL PAGAMENTOS", styles["total_label"]),
                        "",
                        p(pdf_amount(day_data.payment_total), styles["total_value"]),
                        "",
                        "",
                        "",
                        "",
                    ],
                    [
                        p("TOTAL DEVOLUÇÕES", styles["total_label"]),
                        "",
                        p(pdf_amount(day_data.return_total), styles["total_return"]),
                        "",
                        "",
                        "",
                        "",
                    ],
                    [
                        p("TOTAL LÍQUIDO", styles["total_label"]),
                        "",
                        p(pdf_amount(day_data.net_total), styles["total_value"]),
                        "",
                        "",
                        "",
                        "",
                    ],
                    ["", "", "", "", "", "", ""],
                ]
            )
            summary_title_row = len(rows)
            rows.append(
                [
                    p("RESUMO POR MEIO DE PAGAMENTO", styles["summary_title"]),
                    "",
                    "",
                    "",
                    "",
                    "",
                    "",
                ]
            )
            summary_header_row = len(rows)
            rows.append(
                [
                    p("MEIO DE PAGAMENTO", styles["header_left"]),
                    "",
                    "",
                    "",
                    "",
                    "",
                    p("TOTAL", styles["header_right"]),
                ]
            )
            summary_first_row = len(rows)
            rows.extend(
                [
                    [
                        p("Cartão de Crédito", styles["body_left"]),
                        "",
                        "",
                        "",
                        "",
                        "",
                        p("________________ €", styles["body_right"]),
                    ],
                    [
                        p("Dinheiro", styles["body_left"]),
                        "",
                        "",
                        "",
                        "",
                        "",
                        p("________________ €", styles["body_right"]),
                    ],
                    [
                        p("Transferência", styles["body_left"]),
                        "",
                        "",
                        "",
                        "",
                        "",
                        p("________________ €", styles["body_right"]),
                    ],
                    [
                        p("Current Account", styles["body_left"]),
                        "",
                        "",
                        "",
                        "",
                        "",
                        p("Anexar Folha do Programa", styles["current_account"]),
                    ],
                ]
            )

        row_heights = [
            12 * mm,
            8 * mm,
            7 * mm,
            2.5 * mm,
            10 * mm,
        ] + [7.1 * mm] * MIN_MOVEMENT_SLOTS
        if is_final_day_page:
            row_heights += [6.5 * mm, 6.5 * mm, 7 * mm, 3 * mm, 8 * mm, 7 * mm]
            row_heights += [6.5 * mm] * 4

        table = Table(
            rows,
            colWidths=column_widths,
            rowHeights=row_heights,
            repeatRows=0,
        )
        commands: list[tuple[Any, ...]] = [
            ("SPAN", (0, 0), (6, 0)),
            ("SPAN", (0, 1), (6, 1)),
            ("SPAN", (0, 2), (6, 2)),
            ("BACKGROUND", (0, 0), (6, 0), navy),
            ("BACKGROUND", (0, 1), (6, 1), pale_yellow),
            ("BACKGROUND", (0, 4), (6, 4), slate),
            ("GRID", (0, 4), (6, 4 + MIN_MOVEMENT_SLOTS), 0.5, border),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 3),
            ("RIGHTPADDING", (0, 0), (-1, -1), 3),
            ("TOPPADDING", (0, 0), (-1, -1), 1),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 1),
        ]

        movement_table_start = 5
        for slot in range(MIN_MOVEMENT_SLOTS):
            row = movement_table_start + slot
            movement = chunk[slot] if slot < len(chunk) else None
            if movement:
                for column in (1, 4, 6):
                    commands.append(
                        ("BACKGROUND", (column, row), (column, row), pale_yellow)
                    )
                if movement.movement_type == "InvoiceReturn":
                    for column in (0, 2, 3, 5):
                        commands.append(
                            ("BACKGROUND", (column, row), (column, row), pale_red)
                        )
            else:
                for column in (0, 1, 2, 3, 4, 6):
                    commands.append(
                        ("BACKGROUND", (column, row), (column, row), pale_yellow)
                    )
                commands.append(("BACKGROUND", (5, row), (5, row), pale_blue))

        if is_final_day_page:
            assert total_row_start is not None
            assert summary_title_row is not None
            assert summary_header_row is not None
            assert summary_first_row is not None
            for row in range(total_row_start, total_row_start + 3):
                commands.extend(
                    [
                        ("SPAN", (0, row), (1, row)),
                        ("BACKGROUND", (0, row), (2, row), pale_blue),
                        ("GRID", (0, row), (2, row), 0.5, border),
                    ]
                )
            commands.extend(
                [
                    ("LINEABOVE", (0, total_row_start + 2), (2, total_row_start + 2), 1.2, navy),
                    ("LINEBELOW", (0, total_row_start + 2), (2, total_row_start + 2), 1.2, navy),
                    ("SPAN", (0, summary_title_row), (6, summary_title_row)),
                    ("BACKGROUND", (0, summary_title_row), (6, summary_title_row), navy),
                    ("GRID", (0, summary_title_row), (6, summary_title_row), 0.5, border),
                    ("SPAN", (0, summary_header_row), (5, summary_header_row)),
                    ("BACKGROUND", (0, summary_header_row), (6, summary_header_row), slate),
                    ("GRID", (0, summary_header_row), (6, summary_header_row), 0.5, border),
                ]
            )
            for row in range(summary_first_row, summary_first_row + 4):
                commands.extend(
                    [
                        ("SPAN", (0, row), (5, row)),
                        ("GRID", (0, row), (6, row), 0.5, border),
                        ("BACKGROUND", (6, row), (6, row), pale_yellow),
                    ]
                )

        table.setStyle(TableStyle(commands))
        story.append(table)
        if global_index < len(pages) - 1:
            story.append(PageBreak())

    document.build(story)


def summary_payload(
    days: tuple[DayData, ...],
    excel_path: Path,
    pdf_path: Path,
) -> dict[str, Any]:
    return {
        "excel": str(excel_path.resolve()),
        "pdf": str(pdf_path.resolve()),
        "days": [
            {
                "date": day_data.day.isoformat(),
                "provisional": day_data.provisional,
                "movements": len(day_data.movements),
                "invoice_payments": sum(
                    1
                    for movement in day_data.movements
                    if movement.movement_type == "InvoicePayment"
                ),
                "invoice_returns": sum(
                    1
                    for movement in day_data.movements
                    if movement.movement_type == "InvoiceReturn"
                ),
                "payment_total": float(day_data.payment_total),
                "return_total": float(day_data.return_total),
                "net_total": float(day_data.net_total),
            }
            for day_data in days
        ],
    }


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description=(
            "Gera Excel e PDF da folha diária de apoio à caixa a partir "
            "do JSON normalizado."
        )
    )
    parser.add_argument("input", type=Path, help="JSON normalizado.")
    parser.add_argument(
        "--output-dir",
        type=Path,
        required=True,
        help="Diretório para os ficheiros finais.",
    )
    parser.add_argument(
        "--basename",
        help="Nome base opcional, sem caminho e sem extensão.",
    )
    return parser


def main() -> int:
    parser = build_parser()
    args = parser.parse_args()
    try:
        days = load_days(args.input)
        basename = validate_basename(args.basename or default_basename(days))
        args.output_dir.mkdir(parents=True, exist_ok=True)
        excel_path = args.output_dir / f"{basename}.xlsx"
        pdf_path = args.output_dir / f"{basename}.pdf"
        create_excel(days, excel_path)
        create_pdf(days, pdf_path)
        print(
            json.dumps(
                summary_payload(days, excel_path, pdf_path),
                ensure_ascii=False,
                indent=2,
            )
        )
        return 0
    except (InputError, RuntimeError) as exc:
        print(f"ERRO: {exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
