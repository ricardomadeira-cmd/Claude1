#!/usr/bin/env python3
"""Build Hotelbeds portal files from normalized InvoiceXpress data and maintain an audit registry."""

from __future__ import annotations

import argparse
import csv
import json
import sys
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Any

REGISTRY_HEADERS = [
    "event_timestamp",
    "invoice_number",
    "invoicexpress_id",
    "booking_number",
    "original_invoice_date",
    "upload_invoice_date",
    "gross_total_eur",
    "period_from",
    "period_to",
    "batch_filename",
    "status",
    "notes",
]

UPLOAD_HEADERS = [
    "Invoice type",
    "Nº invoice/ credit",
    "Nº to rectify",
    "Invoice date",
    "Booking Nº",
    "Services Description",
    "Costumer Name",
    "Service Date",
    "Nº Days/ Nights",
    "Quantity",
    "Unit price",
    "Currency",
    "Tax",
    "Tax Type",
    "Line total",
    "Invoice total",
]

REVIEW_EXTRA_HEADERS = [
    "Original Invoice Date",
    "Calculated Check-out",
    "Date correction applied",
    "InvoiceXpress ID",
    "Previous registry status",
    "Rounding adjustment",
    "Validation",
]


def q2(value: Decimal) -> Decimal:
    return value.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def dec(value: Any, field: str) -> Decimal:
    if value is None or value == "":
        raise ValueError(f"Missing numeric field: {field}")
    try:
        return Decimal(str(value).replace(" ", "").replace(",", "."))
    except InvalidOperation as exc:
        raise ValueError(f"Invalid decimal in {field}: {value!r}") from exc


def parse_date(value: Any, field: str) -> date:
    if isinstance(value, date) and not isinstance(value, datetime):
        return value
    text = str(value or "").strip()
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    raise ValueError(f"Invalid/missing date in {field}: {value!r}")


def dmy(value: date) -> str:
    return value.strftime("%d-%m-%Y")


def iso(value: date) -> str:
    return value.isoformat()


def money_csv(value: Decimal) -> str:
    return f"{q2(value):.2f}".replace(".", ",")


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def registry_key(row: dict[str, str]) -> tuple[str, str]:
    ix_id = (row.get("invoicexpress_id") or "").strip()
    inv = (row.get("invoice_number") or "").strip()
    return (ix_id, inv)


def ensure_registry(path: Path) -> None:
    if path.exists():
        return
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=REGISTRY_HEADERS, delimiter=";", lineterminator="\r\n")
        writer.writeheader()


def read_registry(path: Path) -> list[dict[str, str]]:
    ensure_registry(path)
    with path.open("r", encoding="utf-8-sig", newline="") as fh:
        reader = csv.DictReader(fh, delimiter=";")
        rows = []
        for row in reader:
            if not any((v or "").strip() for v in row.values()):
                continue
            rows.append({k: (row.get(k) or "") for k in REGISTRY_HEADERS})
        return rows


def append_registry(path: Path, rows: list[dict[str, str]]) -> None:
    ensure_registry(path)
    with path.open("a", encoding="utf-8", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=REGISTRY_HEADERS, delimiter=";", lineterminator="\r\n")
        for row in rows:
            writer.writerow({k: row.get(k, "") for k in REGISTRY_HEADERS})


def latest_registry_map(rows: list[dict[str, str]]) -> dict[tuple[str, str], dict[str, str]]:
    latest: dict[tuple[str, str], dict[str, str]] = {}
    for row in rows:
        key = registry_key(row)
        if not key[0] and not key[1]:
            continue
        latest[key] = row
    return latest


def find_previous(latest: dict[tuple[str, str], dict[str, str]], ix_id: str, invoice_number: str) -> dict[str, str] | None:
    if ix_id:
        for (rid, rinv), row in latest.items():
            if rid and rid == ix_id:
                return row
    for (rid, rinv), row in latest.items():
        if rinv == invoice_number:
            return row
    return None


@dataclass
class ExportLine:
    invoice_number: str
    invoicexpress_id: str
    invoice_type: str
    rectifies: str
    original_invoice_date: date
    upload_invoice_date: date
    booking_number: str
    customer_name: str
    service_date: date
    nights: int
    quantity: Decimal
    gross_line: Decimal
    gross_unit: Decimal
    currency: str
    tax_rate: Decimal
    invoice_total: Decimal
    checkout: date
    previous_status: str
    rounding_adjustment: Decimal = Decimal("0.00")
    date_correction: str = ""

    def upload_row(self) -> list[str]:
        tax = self.tax_rate.normalize()
        tax_text = format(tax, "f")
        return [
            self.invoice_type,
            self.invoice_number,
            self.rectifies,
            dmy(self.upload_invoice_date),
            self.booking_number,
            "Accommodation",
            self.customer_name,
            dmy(self.service_date),
            str(self.nights),
            format(self.quantity.normalize(), "f"),
            money_csv(self.gross_unit),
            self.currency,
            tax_text.replace(".", ","),
            "VAT",
            money_csv(self.gross_line),
            money_csv(self.invoice_total),
        ]

    def review_row(self) -> list[Any]:
        return self.upload_row() + [
            dmy(self.original_invoice_date),
            dmy(self.checkout),
            self.date_correction,
            self.invoicexpress_id,
            self.previous_status,
            money_csv(self.rounding_adjustment) if self.rounding_adjustment else "",
            "OK",
        ]


def load_normalized(path: Path) -> dict[str, Any]:
    with path.open("r", encoding="utf-8-sig") as fh:
        data = json.load(fh)
    if not isinstance(data, dict) or not isinstance(data.get("invoices"), list):
        raise ValueError("Normalized JSON must be an object with an 'invoices' list")
    return data


def build(args: argparse.Namespace) -> int:
    input_path = Path(args.input)
    output_dir = Path(args.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    registry_path = Path(args.registry)

    data = load_normalized(input_path)
    period_from = parse_date(data.get("period_from"), "period_from")
    period_to = parse_date(data.get("period_to"), "period_to")
    if period_to < period_from:
        raise ValueError("period_to cannot be earlier than period_from")

    registry_rows = read_registry(registry_path)
    latest = latest_registry_map(registry_rows)

    exported_invoices: list[dict[str, Any]] = []
    excluded_uploaded: list[str] = []
    prepared_conflicts: list[str] = []
    previous_by_invoice: dict[str, dict[str, str] | None] = {}

    seen_numbers: set[str] = set()
    seen_ids: set[str] = set()

    for inv in data["invoices"]:
        invoice_number = str(inv.get("invoice_number") or "").strip()
        ix_id = str(inv.get("invoicexpress_id") or "").strip()
        if not invoice_number:
            raise ValueError("Invoice without invoice_number")
        if invoice_number in seen_numbers or (ix_id and ix_id in seen_ids):
            raise ValueError(f"Duplicate invoice in normalized input: {invoice_number}")
        seen_numbers.add(invoice_number)
        if ix_id:
            seen_ids.add(ix_id)

        previous = find_previous(latest, ix_id, invoice_number)
        previous_by_invoice[invoice_number] = previous
        status = (previous or {}).get("status", "").strip().lower()
        if status == "uploaded":
            excluded_uploaded.append(invoice_number)
            continue
        if status == "prepared" and not args.reinclude_prepared:
            prepared_conflicts.append(invoice_number)
            continue
        exported_invoices.append(inv)

    if prepared_conflicts and not args.reinclude_prepared:
        conflict_text = ", ".join(prepared_conflicts)
        raise RuntimeError(
            "Invoices already marked prepared require user confirmation before re-export: " + conflict_text
        )

    lines: list[ExportLine] = []
    invoice_notes: dict[str, list[str]] = defaultdict(list)

    for inv in exported_invoices:
        invoice_number = str(inv.get("invoice_number") or "").strip()
        ix_id = str(inv.get("invoicexpress_id") or "").strip()
        doc_type = str(inv.get("document_type") or "invoice").strip().lower()
        invoice_type = "FR" if doc_type in {"credit_note", "credit", "refund"} else "F"
        rectifies = str(inv.get("rectifies_invoice_number") or "").strip()
        original_date = parse_date(inv.get("invoice_date"), f"{invoice_number}.invoice_date")
        booking = str(inv.get("booking_number") or "").strip()
        customer = str(inv.get("customer_name") or "").strip()
        currency = str(inv.get("currency") or "EUR").strip().upper()
        gross_total = q2(dec(inv.get("gross_total"), f"{invoice_number}.gross_total"))
        if not booking:
            raise ValueError(f"{invoice_number}: missing booking_number")
        if not customer:
            raise ValueError(f"{invoice_number}: missing customer_name")
        if currency != "EUR":
            raise ValueError(f"{invoice_number}: unsupported currency {currency}; expected EUR")

        raw_lines = inv.get("lines")
        if not isinstance(raw_lines, list) or not raw_lines:
            raise ValueError(f"{invoice_number}: no accommodation lines")

        prepared_line_data: list[dict[str, Any]] = []
        latest_checkout = original_date
        for idx, item in enumerate(raw_lines, start=1):
            service_date = parse_date(item.get("service_date"), f"{invoice_number}.lines[{idx}].service_date")
            nights = int(item.get("nights") or 1)
            if nights <= 0:
                raise ValueError(f"{invoice_number}: nights must be > 0")
            quantity = dec(item.get("quantity", 1), f"{invoice_number}.lines[{idx}].quantity")
            if quantity <= 0:
                raise ValueError(f"{invoice_number}: quantity must be > 0")
            tax_rate = dec(item.get("tax_rate", 6), f"{invoice_number}.lines[{idx}].tax_rate")
            if item.get("gross_amount") not in (None, ""):
                gross_line = q2(dec(item.get("gross_amount"), f"{invoice_number}.lines[{idx}].gross_amount"))
            else:
                net = dec(item.get("net_amount"), f"{invoice_number}.lines[{idx}].net_amount")
                gross_line = q2(net * (Decimal("1") + tax_rate / Decimal("100")))
            checkout = service_date + timedelta(days=nights)
            latest_checkout = max(latest_checkout, checkout)
            prepared_line_data.append(
                {
                    "service_date": service_date,
                    "nights": nights,
                    "quantity": quantity,
                    "tax_rate": tax_rate,
                    "gross_line": gross_line,
                    "checkout": checkout,
                }
            )

        gross_sum = q2(sum((x["gross_line"] for x in prepared_line_data), Decimal("0.00")))
        diff = q2(gross_total - gross_sum)
        if diff:
            if abs(diff) > Decimal("0.05"):
                raise ValueError(
                    f"{invoice_number}: gross line sum {gross_sum} differs from invoice gross total {gross_total} by {diff}; manual review required"
                )
            prepared_line_data[-1]["gross_line"] = q2(prepared_line_data[-1]["gross_line"] + diff)
            prepared_line_data[-1]["rounding_adjustment"] = diff
            invoice_notes[invoice_number].append(f"Rounding adjustment on last line: {diff:+.2f} EUR")

        upload_date = max(original_date, latest_checkout)
        date_correction = ""
        if upload_date != original_date:
            date_correction = f"{dmy(original_date)} -> {dmy(upload_date)}"
            invoice_notes[invoice_number].append(f"Invoice date adjusted for Hotelbeds checkout rule: {date_correction}")

        previous = previous_by_invoice.get(invoice_number) or {}
        previous_status = previous.get("status", "")

        for x in prepared_line_data:
            gross_line = q2(x["gross_line"])
            gross_unit = q2(gross_line / x["quantity"])
            lines.append(
                ExportLine(
                    invoice_number=invoice_number,
                    invoicexpress_id=ix_id,
                    invoice_type=invoice_type,
                    rectifies=rectifies,
                    original_invoice_date=original_date,
                    upload_invoice_date=upload_date,
                    booking_number=booking,
                    customer_name=customer,
                    service_date=x["service_date"],
                    nights=x["nights"],
                    quantity=x["quantity"],
                    gross_line=gross_line,
                    gross_unit=gross_unit,
                    currency=currency,
                    tax_rate=x["tax_rate"],
                    invoice_total=gross_total,
                    checkout=x["checkout"],
                    previous_status=previous_status,
                    rounding_adjustment=x.get("rounding_adjustment", Decimal("0.00")),
                    date_correction=date_correction,
                )
            )

    if not lines:
        raise RuntimeError("No new invoices to export after registry checks")

    prefix = f"hotelbeds_{iso(period_from)}_{iso(period_to)}"
    main_csv = output_dir / f"{prefix}_COM_IVA.csv"
    review_xlsx = output_dir / f"{prefix}_revisao.xlsx"
    validation_txt = output_dir / f"{prefix}_VALIDACAO.txt"

    with main_csv.open("w", encoding="utf-8", newline="") as fh:
        writer = csv.writer(fh, delimiter=";", quoting=csv.QUOTE_MINIMAL, lineterminator="\r\n")
        for line in lines:
            row = line.upload_row()
            if len(row) != 16:
                raise AssertionError("Upload row must contain exactly 16 columns")
            writer.writerow(row)

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise RuntimeError("openpyxl is required to create the review XLSX") from exc

    wb = Workbook()
    ws = wb.active
    ws.title = "Revisao"
    ws.sheet_view.showGridLines = False
    headers = UPLOAD_HEADERS + REVIEW_EXTRA_HEADERS
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
    for line in lines:
        ws.append(line.review_row())
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    widths = [14, 20, 18, 14, 16, 22, 30, 14, 16, 10, 12, 10, 8, 10, 12, 14, 18, 20, 22, 20, 22, 18, 14]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    wb.save(review_xlsx)

    by_invoice: dict[str, list[ExportLine]] = defaultdict(list)
    for line in lines:
        by_invoice[line.invoice_number].append(line)

    issues: list[str] = []
    for inv_no, inv_lines in by_invoice.items():
        line_sum = q2(sum((x.gross_line for x in inv_lines), Decimal("0.00")))
        inv_total = inv_lines[0].invoice_total
        if line_sum != inv_total:
            issues.append(f"{inv_no}: line sum {line_sum} != invoice total {inv_total}")
        for x in inv_lines:
            if x.upload_invoice_date < x.checkout:
                issues.append(f"{inv_no}: upload invoice date before checkout")

    total_export = q2(sum((inv_lines[0].invoice_total for inv_lines in by_invoice.values()), Decimal("0.00")))
    corrections = [inv for inv, ls in by_invoice.items() if ls[0].date_correction]
    adjusted = [inv for inv, notes in invoice_notes.items() if any("Rounding adjustment" in n for n in notes)]

    report = [
        "HOTELBEDS UPLOAD VALIDATION",
        "",
        f"Period: {iso(period_from)} to {iso(period_to)}",
        f"Invoices in normalized input: {len(data['invoices'])}",
        f"Invoices exported: {len(by_invoice)}",
        f"Exported lines: {len(lines)}",
        f"Already uploaded / excluded: {len(excluded_uploaded)}" + (f" ({', '.join(excluded_uploaded)})" if excluded_uploaded else ""),
        f"Prepared conflicts: {len(prepared_conflicts)}" + (f" ({', '.join(prepared_conflicts)})" if prepared_conflicts else ""),
        f"Gross total exported (IVA included): {total_export:.2f} EUR",
        f"Date corrections: {len(corrections)}" + (f" ({', '.join(corrections)})" if corrections else ""),
        f"Rounding-adjusted invoices: {len(adjusted)}" + (f" ({', '.join(adjusted)})" if adjusted else ""),
        f"Validation issues: {len(issues)}",
    ]
    if issues:
        report += ["", "ISSUES"] + [f"- {x}" for x in issues]
    if invoice_notes:
        report += ["", "NOTES"]
        for inv_no, notes in sorted(invoice_notes.items()):
            for note in notes:
                report.append(f"- {inv_no}: {note}")
    report += ["", f"Upload CSV: {main_csv.name}", f"Review XLSX: {review_xlsx.name}"]
    validation_txt.write_text("\n".join(report) + "\n", encoding="utf-8")

    if issues:
        raise RuntimeError("Validation issues found; inspect validation report before upload")

    prepared_events: list[dict[str, str]] = []
    for inv_no, inv_lines in sorted(by_invoice.items()):
        first = inv_lines[0]
        note_parts = list(invoice_notes.get(inv_no, []))
        if args.reinclude_prepared and first.previous_status.lower() == "prepared":
            note_parts.append("Re-prepared after user confirmation that prior prepared batch was not uploaded")
        prepared_events.append(
            {
                "event_timestamp": utc_now(),
                "invoice_number": inv_no,
                "invoicexpress_id": first.invoicexpress_id,
                "booking_number": first.booking_number,
                "original_invoice_date": iso(first.original_invoice_date),
                "upload_invoice_date": iso(first.upload_invoice_date),
                "gross_total_eur": f"{first.invoice_total:.2f}",
                "period_from": iso(period_from),
                "period_to": iso(period_to),
                "batch_filename": main_csv.name,
                "status": "prepared",
                "notes": " | ".join(note_parts),
            }
        )
    append_registry(registry_path, prepared_events)

    print(f"Created: {main_csv}")
    print(f"Created: {review_xlsx}")
    print(f"Created: {validation_txt}")
    print(f"Updated registry: {registry_path}")
    print(f"Invoices exported: {len(by_invoice)}; gross total: {total_export:.2f} EUR")
    if excluded_uploaded:
        print("Excluded already uploaded: " + ", ".join(excluded_uploaded))
    return 0


def mark(args: argparse.Namespace) -> int:
    registry_path = Path(args.registry)
    rows = read_registry(registry_path)
    latest = latest_registry_map(rows)
    candidates: list[dict[str, str]] = []

    invoice_filter = set(args.invoice or [])
    batch_filter = args.batch or ""

    for row in latest.values():
        if invoice_filter and row.get("invoice_number") in invoice_filter:
            candidates.append(row)
        elif batch_filter and row.get("batch_filename") == batch_filter:
            candidates.append(row)

    if not candidates:
        raise RuntimeError("No matching registry entries found for mark operation")

    events: list[dict[str, str]] = []
    for row in candidates:
        new_row = dict(row)
        new_row["event_timestamp"] = utc_now()
        new_row["status"] = args.status
        new_row["notes"] = args.notes or f"Status changed from {row.get('status', '')} to {args.status}"
        events.append(new_row)
    append_registry(registry_path, events)
    print(f"Marked {len(events)} invoice(s) as {args.status} in {registry_path}")
    return 0


def main() -> int:
    parser = argparse.ArgumentParser(description="Hotelbeds upload builder and duplicate-control registry")
    sub = parser.add_subparsers(dest="command", required=True)

    p_build = sub.add_parser("build", help="Build Hotelbeds upload files from normalized InvoiceXpress JSON")
    p_build.add_argument("--input", required=True, help="Normalized JSON input")
    p_build.add_argument("--registry", required=True, help="Persistent hotelbeds_upload_registry.csv")
    p_build.add_argument("--output-dir", default=".")
    p_build.add_argument("--reinclude-prepared", action="store_true", help="Use only after user confirms prior prepared invoices were not uploaded")
    p_build.set_defaults(func=build)

    p_mark = sub.add_parser("mark", help="Append uploaded/rejected status events to the registry")
    p_mark.add_argument("--registry", required=True)
    p_mark.add_argument("--status", choices=["uploaded", "rejected"], required=True)
    group = p_mark.add_mutually_exclusive_group(required=True)
    group.add_argument("--batch", help="Batch CSV filename")
    group.add_argument("--invoice", nargs="+", help="One or more invoice numbers")
    p_mark.add_argument("--notes", default="")
    p_mark.set_defaults(func=mark)

    args = parser.parse_args()
    try:
        return args.func(args)
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
